"""DOCX 文件格式、并发和缓存生命周期回归。"""

import asyncio
import os
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pytest
from astrbot_plugin_setu.src.infrastructure.sending.docx_service import R18DocxService
from astrbot_plugin_setu.src.shared.send_cache import DiskSendCache
from docx import Document
from PIL import Image, UnidentifiedImageError


def image_bytes(format: str, size=(80, 40), color="blue") -> bytes:
    stream = BytesIO()
    Image.new(
        "RGBA" if isinstance(color, tuple) and len(color) == 4 else "RGB", size, color
    ).save(stream, format=format)
    return stream.getvalue()


def test_docx_contains_all_formats_in_order_and_fits_page(tmp_path):
    images = [
        image_bytes("JPEG", color="red"),
        image_bytes("PNG", (40, 80), color="green"),
        image_bytes("WEBP", (100, 20)),
    ]
    path = R18DocxService(tmp_path).create_docx_with_images(images, ["test"])
    doc = Document(path)
    shapes = list(doc.inline_shapes)
    assert len(shapes) == 3
    assert [round(shape.width / shape.height, 1) for shape in shapes] == [2, 0.5, 5]
    pixels = []
    for shape in shapes:
        rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        with Image.open(BytesIO(doc.part.related_parts[rid].blob)) as embedded:
            pixels.append(embedded.convert("RGB").getpixel((0, 0)))
    expected = []
    for data in images:
        with Image.open(BytesIO(data)) as original:
            expected.append(original.convert("RGB").getpixel((0, 0)))
    assert pixels == expected
    section = doc.sections[0]
    assert all(
        shape.width <= section.page_width - section.left_margin - section.right_margin
        and shape.height
        <= section.page_height - section.top_margin - section.bottom_margin
        for shape in shapes
    )


def test_docx_concurrent_unique_paths(tmp_path):
    service = R18DocxService(tmp_path)
    with ThreadPoolExecutor(max_workers=4) as pool:
        paths = list(
            pool.map(
                lambda _: service.create_docx_with_images([image_bytes("PNG")], []),
                range(4),
            )
        )
    assert len(set(paths)) == 4
    assert all(len(Document(path).inline_shapes) == 1 for path in paths)


@pytest.mark.parametrize("images", [[], [b"invalid"]])
def test_docx_invalid_input_leaves_no_file(tmp_path, images):
    with pytest.raises((ValueError, UnidentifiedImageError)):
        R18DocxService(tmp_path).create_docx_with_images(images, [])
    assert not list(tmp_path.rglob("*.*"))


def test_docx_save_failure_cleans_partial(tmp_path):
    def fail_save(_document, path):
        Path(path).write_bytes(b"partial")
        raise OSError("disk full")

    with (
        patch("docx.document.Document.save", fail_save),
        pytest.raises(OSError, match="disk full"),
    ):
        R18DocxService(tmp_path).create_docx_with_images([image_bytes("PNG")], [])
    assert not list(tmp_path.rglob("*.part"))


@pytest.mark.asyncio
async def test_docx_cleanup_with_image_cache_disabled(tmp_path):
    cache = DiskSendCache(
        tmp_path / "send_cache", enabled=False, ttl_hours=1, max_items=1
    )
    service = R18DocxService(tmp_path)
    paths = await asyncio.gather(
        *[
            asyncio.to_thread(service.create_docx_with_images, [image_bytes("PNG")], [])
            for _ in range(2)
        ]
    )
    assert await cache.cleanup() == 0
    assert all(path.exists() for path in paths)
    for path in paths:
        os.utime(path, (1, 1))
    assert await cache.cleanup() == 2
    assert all(not path.exists() for path in paths)


def test_docx_preserves_webp_transparency(tmp_path):
    path = R18DocxService(tmp_path).create_docx_with_images(
        [image_bytes("WEBP", color=(255, 0, 0, 64))], []
    )
    doc = Document(path)
    shape = doc.inline_shapes[0]
    rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    with Image.open(BytesIO(doc.part.related_parts[rid].blob)) as embedded:
        assert embedded.convert("RGBA").getpixel((0, 0))[3] == 64
